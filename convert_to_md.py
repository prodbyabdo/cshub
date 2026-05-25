import argparse
import os
import fitz  # PyMuPDF
from pptx import Presentation

def extract_pdf(pdf_path):
    try:
        doc = fitz.open(pdf_path)
        text = []
        for i, page in enumerate(doc):
            t = page.get_text()
            if t.strip():
                text.append(f"## Page {i+1}\n\n{t.strip()}")
        return "\n\n---\n\n".join(text)
    except Exception as e:
        print(f"Error reading PDF {pdf_path}: {e}")
        return ""

def extract_pptx(pptx_path):
    try:
        prs = Presentation(pptx_path)
        text = []
        for i, slide in enumerate(prs.slides):
            slide_text = [f"## Slide {i+1}\n"]
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text.strip())
            text.append("\n".join(slide_text))
        return "\n\n---\n\n".join(text)
    except Exception as e:
        print(f"Error reading PPTX {pptx_path}: {e}")
        return ""

def main():
    parser = argparse.ArgumentParser(description="Convert PDF/PPTX to Markdown")
    parser.add_argument("input", help="Input file path (.pdf or .pptx)")
    parser.add_argument("--output", "-o", help="Output Markdown file path (optional)")
    args = parser.parse_args()

    if not os.path.exists(args.input):
        print(f"Error: File '{args.input}' does not exist.")
        return

    ext = os.path.splitext(args.input)[1].lower()
    
    print(f"Processing {args.input}...")
    if ext == ".pdf":
        md_text = extract_pdf(args.input)
    elif ext == ".pptx":
        md_text = extract_pptx(args.input)
    else:
        print(f"Unsupported file extension: {ext}. Only .pdf and .pptx are supported.")
        return

    if not md_text:
        print("No text could be extracted.")
        return

    out_path = args.output
    if not out_path:
        out_path = os.path.splitext(args.input)[0] + ".md"
        
    try:
        with open(out_path, "w", encoding="utf-8") as f:
            f.write(md_text)
        print(f"✓ Successfully converted to {out_path}")
    except Exception as e:
        print(f"Error writing to {out_path}: {e}")

if __name__ == "__main__":
    main()
